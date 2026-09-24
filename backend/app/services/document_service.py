"""
services/document_service.py — upload + text extraction for documents
======================================================================

Supported (deliberately WITHOUT OCR — scanned PDFs are honestly rejected
instead of badly processed; see PLAN/discussion):
    * PDF  (pypdf)          — text-based PDFs, page by page
    * DOCX (python-docx)    — paragraphs + table cells
    * TXT / MD / CSV        — decoded directly

The extraction is deliberately tolerant: an error on page 42 should not
kill the whole document — broken pages are skipped and counted.
"""

import io
import logging
import uuid

from fastapi import UploadFile
from sqlalchemy.ext.asyncio import AsyncSession
from sqlmodel import select

from app.models.document import MAX_TEXT_CHARS, Document
from app.models.research_project import ResearchProject

logger = logging.getLogger(__name__)

# Upload limit: 50 MB — EPUB-converted PDFs like to bring image ballast
# along (30 MB+ with well-selectable text). The EXTRACTED text is capped at
# 1 MB characters anyway; the raw data never lands in the DB.
MAX_UPLOAD_BYTES = 50 * 1024 * 1024

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".csv"}


class DocumentError(Exception):
    """Error during upload/extraction (the endpoint translates it into 4xx)."""


# ============================================================================
# 1) TEXT EXTRACTION per file type
# ============================================================================
def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages: list[str] = []
    broken = 0
    for i, page in enumerate(reader.pages):
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001 — skip a single broken page
            broken += 1
            logger.info("PDF-Seite %d nicht extrahierbar", i + 1)
    if broken:
        logger.warning("PDF: %d/%d Seiten übersprungen", broken, len(reader.pages))
    return "\n\n".join(pages)


def _extract_docx(data: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:  # flatten tables out (cells as rows)
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_plain(data: bytes) -> str:
    # utf-8 first; older German files are often latin-1
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("latin-1", errors="replace")


def extract_text(filename: str, data: bytes) -> str:
    """Extracts text based on the extension and caps it at MAX_TEXT_CHARS."""
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    try:
        if ext == ".pdf":
            text = _extract_pdf(data)
        elif ext == ".docx":
            text = _extract_docx(data)
        elif ext in (".txt", ".md", ".csv"):
            text = _extract_plain(data)
        else:
            raise DocumentError(
                f"Dateityp nicht unterstützt: {ext} (erlaubt: PDF, DOCX, TXT, MD, CSV)"
            )
    except DocumentError:
        raise
    except Exception as e:  # noqa: BLE001 — parser crash -> clean 422 with a reason
        raise DocumentError(
            f"Datei konnte nicht gelesen werden ({type(e).__name__}): {e}"
        ) from e

    if not text.strip():
        # Typical: scanned PDF without a text layer -> be honest instead of 0 characters
        raise DocumentError(
            "Kein Text gefunden — handelt es sich um einen Scan? "
            "OCR (texterkennung) wird nicht unterstützt."
        )
    return text[:MAX_TEXT_CHARS]


# ============================================================================
# 2) CRUD (always ownership-filtered via get_project)
# ============================================================================
async def upload_document(
    session: AsyncSession,
    project: ResearchProject,
    upload: UploadFile,
) -> Document:
    """Reads the file, extracts text and stores the Document."""
    data = await upload.read()
    if len(data) > MAX_UPLOAD_BYTES:
        raise DocumentError(f"Datei zu groß ({len(data) // (1024 * 1024)} MB, max. 50 MB).")
    if not data:
        raise DocumentError("Leere Datei.")

    filename = upload.filename or "datei"
    text = extract_text(filename, data)

    document = Document(
        project_id=project.id,
        filename=filename,
        mime_type=upload.content_type or "application/octet-stream",
        size_bytes=len(data),
        char_count=len(text),
        extracted_text=text,
    )
    session.add(document)
    await session.commit()
    await session.refresh(document)
    return document


async def list_documents(session: AsyncSession, project: ResearchProject) -> list[Document]:
    result = await session.exec(
        select(Document).where(Document.project_id == project.id).order_by(Document.created_at)
    )
    return list(result.all())


async def delete_document(
    session: AsyncSession, project: ResearchProject, document_id: uuid.UUID  # type: ignore[name-defined]
) -> None:
    result = await session.exec(
        select(Document).where(Document.project_id == project.id, Document.id == document_id)
    )
    document = result.first()
    if document is None:
        raise DocumentError("Dokument nicht gefunden.")
    await session.delete(document)
    await session.commit()


async def list_documents_for_chain(
    session: AsyncSession, project: ResearchProject
) -> list[Document]:
    """
    Documents of the ENTIRE conversation: the project's own plus all its
    ancestors (parent_id chain). This keeps uploaded files available across
    follow-up questions — the chat never "loses" them (user request).
    """
    all_docs: list[Document] = []
    seen_projects: set = set()
    current: ResearchProject | None = project
    while current is not None and current.id not in seen_projects:
        seen_projects.add(current.id)
        all_docs.extend(await list_documents(session, current))
        current = (
            await session.get(ResearchProject, current.parent_id)
            if current.parent_id
            else None
        )
    # Oldest first (order = upload chronology of the conversation)
    all_docs.sort(key=lambda d: d.created_at)
    return all_docs


def documents_for_agent(session_documents: list[Document]) -> list[dict]:
    """Compact form for the graph state (name + capped text)."""
    return [
        {"name": d.filename, "text": d.extracted_text[:60_000]}
        for d in session_documents
    ]
